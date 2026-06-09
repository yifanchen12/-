from __future__ import annotations

import sqlite3
from pathlib import Path

from openpyxl import Workbook


EXCEL_MAX_ROWS = 1_048_576


BOOK_COLUMNS = [
    "rec_ctrl_id",
    "top_category_code",
    "top_category_name",
    "category_code",
    "category_name",
    "title",
    "authors",
    "isbn_issn",
    "publisher",
    "publish_date",
    "classno_abs",
    "search_no",
    "subject_terms",
    "library_name",
    "holdings_count",
    "available_count",
    "rating",
    "comment_count",
    "material_type",
    "detail_url",
]

HOLDING_COLUMNS = [
    "rec_ctrl_id",
    "top_category_code",
    "category_code",
    "department",
    "barcode",
    "search_no",
    "register_no",
    "circulation_status",
    "shelf_status",
    "location_url",
    "raw_text",
]


def safe_sheet_name(name: str) -> str:
    invalid_chars = set('[]:*?/\\')
    cleaned = "".join("_" if ch in invalid_chars else ch for ch in name)
    return cleaned[:31] or "Sheet"


def export_excel(db_path: Path, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row

    wb = Workbook(write_only=True)
    summary = wb.create_sheet("Summary")
    summary.append(["top_category_code", "top_category_name", "books"])
    for row in conn.execute(
        """
        SELECT top_category_code, top_category_name, COUNT(*) AS books
        FROM books
        GROUP BY top_category_code, top_category_name
        ORDER BY top_category_code
        """
    ):
        summary.append([row["top_category_code"], row["top_category_name"], row["books"]])

    all_books = wb.create_sheet("All Books")
    all_books.append(BOOK_COLUMNS)
    for row in conn.execute(f"SELECT {', '.join(BOOK_COLUMNS)} FROM books ORDER BY top_category_code, category_code, title"):
        all_books.append([row[column] for column in BOOK_COLUMNS])

    all_holdings = wb.create_sheet("Holdings")
    all_holdings.append(HOLDING_COLUMNS)
    for row in conn.execute(f"SELECT {', '.join(HOLDING_COLUMNS)} FROM holdings ORDER BY top_category_code, category_code, rec_ctrl_id"):
        all_holdings.append([row[column] for column in HOLDING_COLUMNS])

    for row in conn.execute("SELECT DISTINCT top_category_code, top_category_name FROM books ORDER BY top_category_code"):
        sheet = wb.create_sheet(safe_sheet_name(f"{row['top_category_code']} {row['top_category_name']}"))
        sheet.append(BOOK_COLUMNS)
        count = 1
        for book in conn.execute(
            f"SELECT {', '.join(BOOK_COLUMNS)} FROM books WHERE top_category_code = ? ORDER BY category_code, title",
            (row["top_category_code"],),
        ):
            if count >= EXCEL_MAX_ROWS:
                break
            sheet.append([book[column] for column in BOOK_COLUMNS])
            count += 1

    wb.save(output_path)
    conn.close()
